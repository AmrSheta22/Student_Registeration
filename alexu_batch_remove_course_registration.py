"""Remove marked XLSX subjects or one selected registered course from a CSV.

Workflow:
1. Prompt for portal credentials and the student program.
2. Open one visible, detached Chrome session.
3. Find the first valid student in the input CSV and list that student's
   available courses.
4. Prompt for one course and attempt it for every valid student.
5. Checkpoint ``course_removal_status`` with Pandas after every row.
6. Log out normally while leaving the Chrome tab open.

Credentials are kept in memory only. The script never calls ``close()`` or
``quit()``.
"""

from __future__ import annotations

import argparse
import os
import sys
import time
import re
from difflib import SequenceMatcher
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from selenium import webdriver
from selenium.common.exceptions import TimeoutException, StaleElementReferenceException, WebDriverException
from selenium.webdriver import ChromeOptions
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

from student_completed_hours_pipeline import (
    ACADEMIC_YEAR,
    PORTAL_URL,
    PROGRAMS,
    SEMESTER,
    click_select2_option,
    log,
    logout_if_possible,
    resolve_credentials,
    resolve_program,
    xpath_literal,
    normalize_label,
    DomWait,
)


STATUS_COLUMN = "course_removal_status"
STATUS_PENDING = "pending"
STATUS_SUCCEEDED = "succeeded"
STATUS_FAILED = "failed"
WAIT_SECONDS = 180
FAILURE_SCREENSHOT = Path("alexu_batch_removal_failure.png")


def wait_for_dom(driver: webdriver.Chrome, timeout: int = WAIT_SECONDS) -> None:
    WebDriverWait(driver, timeout).until(
        lambda current: current.execute_script("return document.readyState")
        == "complete"
    )


@dataclass(frozen=True)
class Course:
    value: str
    text: str


class CsvBatch:
    """Pandas-backed input validation and per-row status checkpointing."""

    def __init__(self, csv_path: Path) -> None:
        self.path = csv_path.resolve()
        if not self.path.is_file():
            raise ValueError(f"CSV file not found: {self.path}")

        self.dataframe = pd.read_csv(
            self.path,
            dtype={"student_id": "string"},
        )
        if "student_id" not in self.dataframe.columns:
            raise ValueError("Missing required CSV column: student_id")

        self.student_codes = (
            self.dataframe["student_id"].fillna("").str.strip()
        )
        if self.student_codes.empty or (self.student_codes == "").any():
            raise ValueError("Every row must have a student_id")
        if self.student_codes.duplicated().any():
            raise ValueError("student_id values must be unique")

        # A new course choice is a new batch, so prior course results do not
        # apply. Pending values also make an interrupted run obvious.
        self.dataframe[STATUS_COLUMN] = STATUS_PENDING
        self.checkpoint()

    def checkpoint(self) -> None:
        temporary_path = self.path.with_suffix(".csv.tmp")
        self.dataframe.to_csv(
            temporary_path,
            index=False,
            encoding="utf-8",
        )
        deadline = time.monotonic() + 30
        while True:
            try:
                os.replace(temporary_path, self.path)
                return
            except PermissionError:
                if time.monotonic() >= deadline:
                    raise
                time.sleep(0.5)

    def set_status(self, row_index: int, status: str) -> None:
        if status not in {STATUS_SUCCEEDED, STATUS_FAILED}:
            raise ValueError(f"Invalid final status: {status}")
        self.dataframe.at[row_index, STATUS_COLUMN] = status
        self.checkpoint()


class PortalSession:
    """Selenium operations for one authenticated registrar session."""

    def __init__(self, timeout: int = WAIT_SECONDS) -> None:
        options = ChromeOptions()
        options.add_experimental_option("detach", True)
        options.add_argument("--start-maximized")
        log("Opening visible detached Chrome; the tab will remain open.")
        self.driver = webdriver.Chrome(options=options)
        self.wait = DomWait(self.driver, timeout)
        self.logged_in = False
        self.stage = 'opening browser'

    def click_navigation(self, locator, stage):
        """Reacquire navigation controls inside the wait, not after it."""
        self.stage = stage
        def click_fresh(current):
            if current.execute_script('return document.readyState') != 'complete':
                return False
            element = EC.element_to_be_clickable(locator)(current)
            if not element:
                return False
            element.click()
            return True
        self.wait.until(click_fresh, f'Navigation failed during {stage}')

    def login_and_choose_scope(
        self,
        username: str,
        password: str,
        program: str,
        year: str = ACADEMIC_YEAR,
        semester: str = SEMESTER,
    ) -> None:
        self.driver.get(PORTAL_URL)
        wait_for_dom(self.driver)
        self.stage = 'filling both credentials and submitting login'
        # Resolve fresh fields, populate, verify and click in one browser task.
        # This prevents a refreshed login page from losing one field between
        # separate Selenium clear/send_keys/click calls. Never return credentials.
        def fill_and_submit(current):
            return current.execute_script("""
                if (document.readyState !== 'complete') return false;
                const visible = element => element && !element.disabled &&
                    element.getClientRects().length > 0 &&
                    getComputedStyle(element).visibility !== 'hidden';
                const user = document.getElementById('username');
                const pass = document.getElementById('password');
                const buttons = [...document.querySelectorAll(
                    "input[type='submit'], button[type='submit']")].filter(visible);
                if (!visible(user) || !visible(pass) || buttons.length !== 1) return false;
                const setter = Object.getOwnPropertyDescriptor(HTMLInputElement.prototype, 'value').set;
                for (const [element, value] of [[user, arguments[0]], [pass, arguments[1]]]) {
                    setter.call(element, value);
                    element.dispatchEvent(new Event('input', {bubbles:true}));
                    element.dispatchEvent(new Event('change', {bubbles:true}));
                }
                if (!user.isConnected || !pass.isConnected || !buttons[0].isConnected ||
                    user.value !== arguments[0] || pass.value !== arguments[1] ||
                    !user.value || !pass.value) return false;
                buttons[0].click();
                return true;
            """, username, password)
        self.wait.until(fill_and_submit)

        self.stage = 'waiting for login transition'
        reports_label = self.wait.until(
            EC.element_to_be_clickable(
                (
                    By.XPATH,
                    "//span[contains(@class,'hide-menu') and "
                    "contains(normalize-space(.),'Student') and contains(normalize-space(.),'Report')]",
                )
            )
        )
        self.logged_in = True
        log("Login succeeded.")

        self.click_navigation((By.XPATH, "//a[.//span[contains(@class,'hide-menu') and contains(normalize-space(.),'Student') and contains(normalize-space(.),'Report')]]"), 'opening reports scope')
        wait_for_dom(self.driver)
        self.wait.until(EC.presence_of_element_located((By.ID, "programBox")))

        self.stage = 'selecting program'
        click_select2_option(
            self.wait,
            "select2-programBox-container",
            "select2-programBox-results",
            program,
        )
        self.stage = 'selecting academic year'
        click_select2_option(
            self.wait,
            "select2-yearBox-container",
            "select2-yearBox-results",
            year,
        )
        self.stage = 'selecting semester'
        click_select2_option(
            self.wait,
            "select2-semesterBox-container",
            "select2-semesterBox-results",
            semester,
        )
        self.click_navigation((By.CSS_SELECTOR, "button[type='submit'][name='Accept']"), 'accepting academic scope')
        wait_for_dom(self.driver)

        self.click_navigation((By.XPATH, "//a[contains(@href,'registrar_remove.php')][.//*[contains(normalize-space(.),'Remove Course Registration') or contains(normalize-space(.),'Remove Registration')]]"), 'opening removal form')
        wait_for_dom(self.driver)
        self.wait.until(EC.visibility_of_element_located((By.ID, "scodeBox")))
        log("Opened the clean student-code form.")

    @staticmethod
    def missing_student_message(student_code: str) -> str:
        return (
            "Error: No Matching Student Record in the Selected Program "
            f"Was Found for {student_code}"
        )

    def open_student(self, student_code: str) -> bool:
        """Return False when the portal refreshes with the missing-ID error."""
        code_box = self.wait.until(
            EC.visibility_of_element_located((By.ID, "scodeBox"))
        )
        code_box.clear()
        code_box.send_keys(student_code)
        self.wait.until(
            EC.element_to_be_clickable(
                (
                    By.CSS_SELECTOR,
                    "form[name='toplistform'] button[type='submit']",
                )
            )
        ).click()

        missing_message = self.missing_student_message(student_code)

        def student_result(current: webdriver.Chrome) -> str | bool:
            if current.find_elements(By.ID, "courseBox"):
                return "valid"
            if missing_message in current.page_source:
                return "missing"
            return False

        result = self.wait.until(student_result)
        wait_for_dom(self.driver)
        if result == "missing":
            # The same clean form is already present after the refresh.
            self.wait.until(
                EC.visibility_of_element_located((By.ID, "scodeBox"))
            )
            return False
        return True

    def available_courses(self) -> list[Course]:
        options = self.wait.until(
            EC.presence_of_all_elements_located(
                (By.CSS_SELECTOR, "#courseBox option")
            )
        )
        courses: list[Course] = []
        for option in options:
            value = (option.get_attribute("value") or "").strip()
            text = option.text.strip()
            if value and text and text != "Select Registered Course to Remove":
                courses.append(Course(value=value, text=text))
        return courses

    def select_course(self, course: Course) -> bool:
        """Select the exact course, or return False if this student lacks it."""
        self.stage = 'opening course selector'
        self.wait.until(lambda current: current.find_element(By.ID, 'courseBox').is_enabled())
        self.wait.until(
            EC.element_to_be_clickable(
                (
                    By.XPATH,
                    "//*[@id='select2-courseBox-container']"
                    "/following-sibling::span"
                    "[contains(@class,'select2-selection__arrow')]",
                )
            )
        ).click()

        self.stage = 'matching course selector option'
        options = self.wait.until(EC.presence_of_all_elements_located((By.CSS_SELECTOR, '#select2-courseBox-results li')))
        matches = [option for option in options if normalize_label(option.text) == normalize_label(course.text)
                   and option.get_attribute('aria-disabled') != 'true']
        if len(matches) != 1:
            self.driver.find_element(By.TAG_NAME, "body").send_keys(Keys.ESCAPE)
            return False
        matches[0].click()
        self.stage = 'verifying selected course value'
        self.wait.until(
            lambda current: current.find_element(By.ID, 'courseBox').get_attribute('value') == course.value
        )
        return True

    def submit_removal(self) -> tuple[bool, str]:
        self.stage = 'waiting for previous error notification to clear'
        # An old error toast must disappear before another submission.
        self.wait.until(lambda current: not any(e.is_displayed() for e in current.find_elements(By.CSS_SELECTOR, 'article.alertify-log-error')))
        # Latch short-lived toasts in the page, independently of Selenium polls.
        self.driver.execute_script("""
            if (window.__registrationErrorsObserver) window.__registrationErrorsObserver.disconnect();
            window.__registrationErrors = [];
            const baseline = new WeakSet(document.querySelectorAll('article.alertify-log-error'));
            const capture = node => {
                if (!(node instanceof Element)) return;
                const articles = [...node.querySelectorAll('article.alertify-log-error')];
                if (node.matches('article.alertify-log-error')) articles.push(node);
                const parent = node.closest('article.alertify-log-error');
                if (parent) articles.push(parent);
                for (const article of articles) {
                    if (baseline.has(article)) continue;
                    const center = article.querySelector('center');
                    const message = (center ? center.textContent : article.textContent).trim();
                    if (message && !window.__registrationErrors.includes(message)) window.__registrationErrors.push(message);
                }
            };
            window.__registrationErrorsObserver = new MutationObserver(records => {
                for (const record of records) {
                    if (record.type === 'attributes' && record.target.matches('article.alertify-log-error.alertify-log-show')) baseline.delete(record.target);
                    capture(record.target.nodeType === 1 ? record.target : record.target.parentElement);
                    for (const node of record.addedNodes) { baseline.delete(node); capture(node); }
                    for (const node of record.removedNodes) capture(node);
                }
            });
            window.__registrationErrorsObserver.observe(document.documentElement,
                {childList:true, subtree:true, characterData:true, attributes:true, attributeFilter:['class']});
        """)
        self.stage = 'clicking removal button'
        self.wait.until(
            EC.element_to_be_clickable(
                (
                    By.CSS_SELECTOR,
                    "button[type='submit'][onclick='deleteRecord();']",
                )
            )
        ).click()
        def confirmation_or_error(current):
            self.check_session()
            messages = current.execute_script('return window.__registrationErrors || []')
            if messages:
                return ('error', messages[-1])
            titles = current.find_elements(By.XPATH, "//button[contains(@class,'swal2-confirm') and normalize-space(.)='Yes, delete it!']")
            if titles and titles[0].is_displayed():
                return ('confirm', '')
            return False
        self.stage = 'waiting for removal confirmation'
        kind, detail = self.wait.until(confirmation_or_error)
        if kind == 'error':
            return False, detail
        self.stage = 'confirming removal submission'
        selected_label = self.driver.execute_script("""
            const box = document.getElementById('courseBox');
            return box && box.selectedIndex >= 0 ? box.options[box.selectedIndex].text : '';
        """)
        if not selected_label:
            raise RuntimeError('No selected course label available before removal confirmation')
        document_token = str(time.time_ns())
        self.driver.execute_script('window.__removalDocumentToken = arguments[0]', document_token)
        self.wait.until(
            EC.element_to_be_clickable(
                (
                    By.XPATH,
                    "//button[contains(@class,'swal2-confirm') and "
                    "normalize-space(.)='Yes, delete it!']",
                )
            )
        ).click()

        def outcome(current: webdriver.Chrome) -> tuple[str, str] | bool:
            self.check_session()
            messages = current.execute_script('return window.__registrationErrors || []')
            if messages:
                return ('error', messages[-1])
            success = current.find_elements(
                By.XPATH,
                "//*[@id='swal2-title' and "
                "(contains(translate(normalize-space(.),'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'deleted') or contains(translate(normalize-space(.),'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'removed'))]",
            )
            success_icons = current.find_elements(By.CSS_SELECTOR, '.swal2-success')
            # The success icon is authoritative; deletion titles vary by portal page.
            if any(icon.is_displayed() for icon in success_icons):
                return ("success", "Course Registration Removed")

            for error in current.find_elements(
                By.CSS_SELECTOR,
                ".alertify-log-error",
            ):
                if error.is_displayed():
                    centers = error.find_elements(By.TAG_NAME, "center")
                    return (
                        "error",
                        (centers[0].text.strip() if centers else error.text.strip()) or "Portal rejected request",
                    )
            # deleteRecord redirects on BOTH AJAX completion and AJAX error.
            # Wait for a new document, then verify the refreshed registered list.
            reloaded = current.execute_script("return document.readyState === 'complete' && window.__removalDocumentToken !== arguments[0]", document_token)
            if reloaded:
                previous = current.find_elements(By.XPATH, "//a[contains(@href,'registrar_remove.php?') and normalize-space(.)='Previous']")
                if any(link.is_displayed() for link in previous):
                    boxes = current.find_elements(By.ID, 'courseBox')
                    if boxes:
                        options = boxes[0].find_elements(By.TAG_NAME, 'option')
                        if any(normalize_label(option.text) == normalize_label(selected_label) for option in options):
                            return ('error', 'course remains registered after deletion request')
                        return ('success', 'selected course no longer appears in the refreshed registered-course list')
                    return ('error', 'outcome unconfirmed: refreshed page has no registered-course list; verify portal before retrying')
            return False

        self.stage = 'waiting for deletion reload and verifying registered courses'
        try:
            kind, detail = self.wait.until(outcome)
        except TimeoutException:
            # Read the retained notification one final time at the deadline.
            messages = self.driver.execute_script('return window.__registrationErrors || []')
            if messages:
                return False, messages[-1]
            raise
        if kind == "error":
            return False, detail

        self.stage = 'dismissing confirmed success'
        def dismiss_success(current):
            if not any(icon.is_displayed() for icon in current.find_elements(By.CSS_SELECTOR, '.swal2-success')):
                return True  # The dialog may dismiss itself.
            for button in current.find_elements(By.CSS_SELECTOR, '.swal2-popup .swal2-confirm'):
                if button.is_displayed() and button.is_enabled():
                    button.click()
                    return True
            return False
        self.wait.until(dismiss_success)
        self.wait.until(lambda current: not any(icon.is_displayed() for icon in current.find_elements(By.CSS_SELECTOR, '.swal2-success')))
        return True, detail

    def check_session(self) -> None:
        if any(element.is_displayed() for element in self.driver.find_elements(By.ID, 'username')):
            raise RuntimeError(f'Portal returned to login page during {self.stage}')

    def return_to_student_form(self) -> None:
        self.stage = 'waiting for removal page to finish loading'
        self.wait.until(lambda current: current.execute_script('return document.readyState') == 'complete')
        def clean_form(current):
            return (any(box.is_displayed() for box in current.find_elements(By.ID, 'scodeBox'))
                    and not any(box.is_displayed() for box in current.find_elements(By.ID, 'courseBox')))
        self.stage = 'locating session-specific Previous link'
        def previous_target(current):
            if current.execute_script('return document.readyState') != 'complete':
                return False
            links = current.find_elements(By.XPATH,
                "//a[contains(@href,'registrar_remove.php?') and normalize-space(.)='Previous']")
            for link in links:
                if link.is_displayed():
                    return link.get_attribute('href')
            return '__already_on_form__' if clean_form(current) else False
        href = self.wait.until(previous_target, 'Removal page has no visible session-specific Previous link')
        if href != '__already_on_form__':
            self.stage = 'following Previous link to student-ID form'
            # Follow only the supplied Previous anchor, never an arbitrary sidebar link.
            self.driver.get(href)
            self.wait.until(lambda current: current.execute_script('return document.readyState') == 'complete')
            self.wait.until(EC.visibility_of_element_located((By.ID, 'scodeBox')))
        log('Returned through Previous to the removal student-ID form.')

    def recover_student_form(self) -> bool:
        if self.driver.find_elements(By.ID, "scodeBox"):
            return True
        previous = self.driver.find_elements(
            By.XPATH,
            "//a[contains(@href,'registrar_remove.php') and "
            "normalize-space(.)='Previous']",
        )
        if previous:
            previous[0].click()
            wait_for_dom(self.driver)
            self.wait.until(
                EC.visibility_of_element_located((By.ID, "scodeBox"))
            )
            return True
        return False

    def logout(self) -> bool:
        try:
            return self._logout()
        except Exception as exc:
            log(f"Logout NOT confirmed: {type(exc).__name__}. Keep the website open and log out manually.")
            return False

    def _logout(self) -> bool:
        # A session may already have expired or have logged out successfully.
        if self.driver.find_elements(By.ID, 'username') and self.driver.find_elements(By.ID, 'password'):
            if (self.driver.find_element(By.ID, 'username').is_displayed()
                    and self.driver.find_element(By.ID, 'password').is_displayed()
                    and not self.driver.find_elements(By.CSS_SELECTOR, 'a.pro-pic')):
                self.logged_in = False
                log('Logout confirmed on the login page; Chrome remains open.')
                return True
        confirmed = False
        authenticated = bool(
            self.driver.find_elements(
                By.XPATH,
                "//a[contains(@class,'pro-pic')]",
            )
        )
        if self.logged_in and not authenticated:
            for _ in range(2):
                self.driver.back()
                try:
                    WebDriverWait(self.driver, 5).until(EC.presence_of_element_located((By.CSS_SELECTOR, 'a.pro-pic')))
                    break
                except TimeoutException:
                    pass
        if self.logged_in or authenticated or self.driver.find_elements(By.CSS_SELECTOR, "a[href*='logout.php']"):
            confirmed = logout_if_possible(
                self.driver,
                DomWait(self.driver, 10),
            )
            if not confirmed:
                log("Logout NOT confirmed. Keep the website open for manual logout.")
            else:
                self.logged_in = False
        time.sleep(2)
        log("Detached Chrome remains open.")
        return confirmed


def choose_course(courses: list[Course]) -> Course:
    if not courses:
        raise ValueError("The first valid student has no available courses")

    print("\nAvailable courses for the first valid student:")
    for number, course in enumerate(courses, start=1):
        print(f"  {number}. {course.text}")

    while True:
        choice = input(f"Course [1-{len(courses)}]: ").strip()
        if choice.isdigit() and 1 <= int(choice) <= len(courses):
            return courses[int(choice) - 1]
        print(f"Please enter a number from 1 to {len(courses)}.")


def course_identity(text: str) -> tuple[str, str, str, str] | None:
    """Never fuzzy-match across course codes, groups, years or terms."""
    code = re.match(r"\s*([\w]+(?:-[\w]+)*)\s*-", text)
    group = re.search(r"Group\s*:\s*(\d+)", text, re.I)
    year = re.search(r"\d{4}\s*/\s*\d{4}", text)
    term = re.search(r"\b(Fall|Summer|Spring|Winter)\b", text, re.I)
    if not all((code, group, year, term)):
        return None
    return (code.group(1).casefold(), group.group(1), normalize_label(year.group()), term.group().casefold())


def match_course(header: str, courses: list[Course]) -> Course | None:
    exact = [c for c in courses if normalize_label(c.text) == normalize_label(header)]
    if len(exact) == 1:
        return exact[0]
    identity = course_identity(header)
    if identity is None:
        return None
    scored = sorted([(SequenceMatcher(None, normalize_label(header), normalize_label(c.text)).ratio(), c)
                     for c in courses if course_identity(c.text) == identity], key=lambda pair: pair[0], reverse=True)
    if scored and scored[0][0] >= .94 and (len(scored) == 1 or scored[0][0] - scored[1][0] >= .03):
        return scored[0][1]
    return None


class WorkbookBatch:
    """Keep input sheets intact and checkpoint subject results in a log sheet."""
    def __init__(self, path: Path, sheet: str | None, limit: int | None, year: str, semester: str,
                 resume: bool = False, start_student: str | None = None, retry_failed: bool = False):
        from openpyxl import load_workbook
        self.path = path.resolve()
        self.book = load_workbook(self.path)
        if sheet and sheet not in self.book.sheetnames:
            raise ValueError(f'Worksheet not found: {sheet}')
        self.source = self.book[sheet] if sheet else self.book.worksheets[0]
        if self.source.title == 'Removal log':
            raise ValueError('Choose an input worksheet, not Removal log')
        self.headers = [str(c.value or '').strip() for c in self.source[1]]
        if normalize_label(self.headers[0]) not in {'id', 'studentid'}:
            raise ValueError('First column must be ID or student_id')
        if any(not h for h in self.headers) or len(set(map(normalize_label, self.headers))) != len(self.headers):
            raise ValueError('Headers must be nonempty and unique')
        if limit is not None and limit < 1:
            raise ValueError('--limit must be positive')
        self.rows = []
        seen = set()
        for row in self.source.iter_rows(min_row=2):
            if all(c.value is None for c in row):
                continue
            value = row[0].value
            code = str(int(value)) if isinstance(value, (float, int)) and float(value).is_integer() else str(value or '').strip()
            if not code or not code.isdigit() or code in seen:
                raise ValueError(f'Invalid or duplicate ID at row {row[0].row}')
            seen.add(code)
            requested = []
            for column, cell in enumerate(row[1:], start=2):
                if cell.value in (None, '', 0, False):
                    continue
                if str(cell.value).strip().casefold() not in {'1', '1.0', 'true', 'yes', 'x'}:
                    raise ValueError(f'Unsupported request marker in {cell.coordinate}: {cell.value!r}')
                identity = course_identity(self.headers[column-1])
                if identity and (identity[2] != normalize_label(year) or identity[3] != semester.casefold()):
                    raise ValueError(f'Course scope differs from selected year/semester: {self.headers[column-1]}')
                requested.append(column)
            self.rows.append((row[0].row, code, requested))
        if start_student:
            positions = [i for i, (_, code, _) in enumerate(self.rows) if code == start_student.strip()]
            if not positions:
                raise ValueError(f'Start student not found: {start_student}')
            self.rows = self.rows[positions[0]:]
        if not self.rows:
            raise ValueError('No students found')
        existing = 'Removal log' in self.book.sheetnames
        if existing and not resume:
            raise ValueError('Removal log already exists; use a fresh workbook copy to prevent accidental repeat submissions')
        if retry_failed and not resume:
            raise ValueError('--retry-failed requires --resume')
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        self.sheet = self.book['Removal log'] if existing else self.book.create_sheet('Removal log')
        if existing:
            logged_headers = [str(c.value or '').strip() for c in self.sheet[1]]
            if logged_headers != self.headers:
                raise ValueError('Log headers differ from input; cannot safely resume')
        else:
            self.sheet.append(self.headers)
        last_timeout = None
        if existing and resume and not retry_failed:
            # Sequential runs follow worksheet order; find the last failed
            # timeout in the selected starting range before applying --limit.
            for position, (row, code, columns) in enumerate(self.rows):
                for column in columns:
                    status = str(self.sheet.cell(row, column).value or '').casefold()
                    if status.startswith('failed') and ('timeoutexception' in status or 'timeout' in status):
                        last_timeout = (row, column)
                        timeout_position = position
            if last_timeout:
                self.rows = self.rows[timeout_position:]
                row, code, columns = self.rows[0]
                self.rows[0] = (row, code, [column for column in columns if column >= last_timeout[1]])
                log(f'Resuming at the last timeout: student {code}, subject {self.headers[last_timeout[1]-1]}. This subject will be retried first.')
        if limit:
            self.rows = self.rows[:limit]
        self.sheet.freeze_panes = 'B2'
        self.sheet.column_dimensions['A'].width = 18
        self.sheet.row_dimensions[1].height = 120
        for column in range(1, len(self.headers) + 1):
            cell = self.sheet.cell(1, column)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='234E70')
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if column > 1:
                self.sheet.column_dimensions[get_column_letter(column)].width = 36
        remaining = []
        for row, code, columns in self.rows:
            logged_id = self.sheet.cell(row, 1).value
            if logged_id is not None and str(logged_id).strip() != code:
                raise ValueError(f'Log student ID differs at row {row}; cannot safely resume')
            self.sheet.cell(row, 1, code)
            pending = []
            for column in columns:
                status = str(self.sheet.cell(row, column).value or '').strip()
                retry_last_timeout = (row, column) == last_timeout
                if 'outcome unconfirmed' in status.casefold() and not retry_last_timeout:
                    if 'skipped without retry' in status.casefold() and not retry_failed:
                        continue
                    raise ValueError(f'Unconfirmed result for student {code}, column {column}. Verify on portal and set the log cell to succeeded, Failed: <reason>, or pending before resuming.')
                if status.casefold() == 'succeeded':
                    continue
                if status.casefold().startswith('failed') and not retry_failed and not retry_last_timeout:
                    continue
                if status and status.casefold() != 'pending' and not status.casefold().startswith('failed'):
                    raise ValueError(f'Unexpected log status for student {code}, column {column}: {status}')
                self.sheet.cell(row, column, 'pending')
                pending.append(column)
            remaining.append((row, code, pending))
        self.rows = remaining
        self.checkpoint()

    def checkpoint(self):
        temporary = self.path.with_name(self.path.stem + '.checkpoint.xlsx')
        self.book.save(temporary)
        deadline = time.monotonic() + 30
        while True:
            try:
                os.replace(temporary, self.path)
                return
            except PermissionError:
                if time.monotonic() >= deadline:
                    raise
                time.sleep(.5)

    def result(self, row, column, status):
        from openpyxl.styles import Alignment, PatternFill
        self.sheet.cell(row, column, status)
        self.sheet.cell(row, column).alignment = Alignment(wrap_text=True, vertical='top')
        self.sheet.cell(row, column).fill = PatternFill('solid', fgColor='E2F0D9' if status == 'succeeded' else 'FCE4D6')
        self.checkpoint()
        log(f"Checkpointed student {self.sheet.cell(row, 1).value}: {self.headers[column-1]} = {status}")


def run_workbook(args, username, password, program):
    batch = WorkbookBatch(args.csv_path, args.sheet, args.limit, args.year, args.semester,
                          args.resume, args.start_student, args.retry_failed)
    if not any(columns for _, _, columns in batch.rows):
        log('No unprocessed subjects in the selected students; no portal session opened.')
        return 0
    portal = None
    def reconnect():
        log('Unexpected portal failure. Confirming logout before reconnecting.')
        try:
            portal.driver.save_screenshot(str(FAILURE_SCREENSHOT.resolve()))
        except Exception:
            pass
        if not portal.logout():
            raise RuntimeError('Logout was not confirmed; reconnect cancelled. Leave the website open for manual logout.')
        # Reuse the same visible browser; credentials stay in memory.
        for attempt in range(1, 4):
            try:
                log(f'Recovery login/scope attempt {attempt}/3.')
                portal.login_and_choose_scope(username, password, program, args.year, args.semester)
                log('Logged in again with the same scope; continuing after the skipped subject.')
                return
            except (TimeoutException, StaleElementReferenceException, WebDriverException) as exc:
                log(f'Recovery failed during {portal.stage}: {type(exc).__name__}.')
                if not portal.logout():
                    raise RuntimeError('Recovery logout not confirmed; stopping with Chrome open.') from exc
                if attempt == 3:
                    raise RuntimeError(f'Recovery exhausted after 3 attempts at {portal.stage}') from exc
                time.sleep(2)
    try:
        portal = PortalSession()
        portal.login_and_choose_scope(username, password, program, args.year, args.semester)
        for row, code, columns in batch.rows:
            if not columns:
                continue
            student_open = False
            for column in columns:
                header = batch.headers[column-1]
                status = None
                failure = None
                missing_student = False
                try:
                    if not student_open:
                        if not portal.open_student(code):
                            missing_student = True
                            status = 'Failed: student not found in selected program'
                        else:
                            student_open = True
                    if not missing_student:
                        course = match_course(header, portal.available_courses())
                        if course is None or not portal.select_course(course):
                            status = 'Failed: not registered or unavailable for removal'
                        else:
                            log(f'Student {code}: submitting {course.text}')
                            succeeded, detail = portal.submit_removal()
                            status = 'succeeded' if succeeded else f'Failed: {detail}'
                        portal.return_to_student_form()
                        student_open = False
                except Exception as exc:
                    failure = exc
                    # Keep an already confirmed outcome if subsequent page recovery failed.
                    if status is None:
                        status = f'Failed: skipped without retry; outcome unconfirmed (stage: {portal.stage}; {type(exc).__name__}: {exc})'
                # Checkpoint errors must stop the run rather than trigger more submissions.
                if missing_student:
                    for remaining in columns[columns.index(column):]:
                        batch.result(row, remaining, status)
                    break
                batch.result(row, column, status)
                if failure is not None:
                    log(f'Student {code}: skipping affected subject after {type(failure).__name__}.')
                    reconnect()
                    student_open = False
            if student_open:
                try:
                    portal.return_to_student_form()
                except Exception:
                    reconnect()
        return 0
    except (KeyboardInterrupt, EOFError):
        log('Interrupted; completed subject checkpoints retained.')
        return 130
    except Exception as exc:
        log(f'Workbook removal stopped: {type(exc).__name__}: {exc}')
        if portal:
            try:
                portal.driver.save_screenshot(str(FAILURE_SCREENSHOT.resolve()))
            except Exception:
                pass
        return 1
    finally:
        if portal:
            portal.logout()


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Remove requested XLSX subjects per student, or choose one registered course for a legacy CSV batch."
    )
    parser.add_argument(
        "csv_path",
        nargs="?",
        type=Path,
        default=Path("data.csv"),
        help="Input CSV or XLSX updated in place (default: data.csv).",
    )
    parser.add_argument(
        "--username",
        help="Portal username; prompts when omitted.",
    )
    parser.add_argument(
        "--program",
        choices=PROGRAMS,
        help="Student program; shows the existing numbered menu when omitted.",
    )
    parser.add_argument('--year', default=ACADEMIC_YEAR)
    parser.add_argument('--semester', default=SEMESTER)
    parser.add_argument('--sheet', help='Input workbook worksheet name')
    parser.add_argument('--limit', type=int, help='Process only the first N students')
    parser.add_argument('--resume', action='store_true', help='Preserve successes; restart at the last timeout, or process blank/pending requests if none')
    parser.add_argument('--start-student', help='Start at this student ID, inclusively; --limit applies from here')
    parser.add_argument('--retry-failed', action='store_true', help='With --resume, also retry confirmed failures; never repeat successes')
    return parser.parse_args()


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="backslashreplace")

    args = parse_arguments()
    try:
        # Reuse the exact credential and program resolution logic from the
        # completed-hours pipeline.
        username, password = resolve_credentials(args)
        program = resolve_program(args, interactive=True)
        if args.csv_path.suffix.lower() == '.xlsx':
            return run_workbook(args, username, password, program)
        batch = CsvBatch(args.csv_path)
    except ValueError as exc:
        print(str(exc), file=sys.stderr)
        return 2

    portal = PortalSession()
    chosen_course: Course | None = None
    try:
        portal.login_and_choose_scope(username, password, program, args.year, args.semester)
        log(f"Loaded {len(batch.dataframe)} students from {batch.path}.")

        for position, (row_index, student_code) in enumerate(
            batch.student_codes.items(),
            start=1,
        ):
            log(
                f"[{position}/{len(batch.dataframe)}] "
                f"Checking student {student_code}."
            )
            status = STATUS_FAILED
            try:
                if not portal.open_student(student_code):
                    log(
                        f"[{position}/{len(batch.dataframe)}] "
                        "No matching student record; skipped."
                    )
                    continue

                if chosen_course is None:
                    chosen_course = choose_course(portal.available_courses())
                    log(f"Selected batch course: {chosen_course.text!r}.")

                if not portal.select_course(chosen_course):
                    log(
                        f"[{position}/{len(batch.dataframe)}] "
                        "Selected course is unavailable for this student."
                    )
                else:
                    succeeded, detail = portal.submit_removal()
                    if succeeded:
                        status = STATUS_SUCCEEDED
                        log(
                            f"[{position}/{len(batch.dataframe)}] "
                            "Removal succeeded."
                        )
                    else:
                        log(
                            f"[{position}/{len(batch.dataframe)}] "
                            f"Removal failed: {detail}"
                        )
                portal.return_to_student_form()
            except Exception as exc:
                log(
                    f"[{position}/{len(batch.dataframe)}] "
                    f"Removal failed: {type(exc).__name__}: {exc}"
                )
                if not portal.recover_student_form():
                    raise RuntimeError(
                        "Could not recover the clean student-code form"
                    ) from exc
            finally:
                batch.set_status(row_index, status)
                log(
                    f"[{position}/{len(batch.dataframe)}] "
                    f"Checkpointed {STATUS_COLUMN}={status}."
                )

        if chosen_course is None:
            log("No valid student exposed a registered-course list; nothing was removed.")
        return 0
    except (KeyboardInterrupt, EOFError):
        log("Interactive input was cancelled; checkpointed results were kept.")
        return 130
    except Exception as exc:
        log(f"Batch workflow failed: {type(exc).__name__}: {exc}")
        try:
            portal.driver.save_screenshot(str(FAILURE_SCREENSHOT.resolve()))
            log(f"Saved diagnostic screenshot to {FAILURE_SCREENSHOT.resolve()}.")
        except Exception:
            pass
        return 1
    finally:
        portal.logout()


if __name__ == "__main__":
    raise SystemExit(main())


